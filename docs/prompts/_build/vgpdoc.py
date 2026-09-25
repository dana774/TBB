"""VGP house-style docx builder (Linux/python-docx port of md2docx.ps1).

House style: Segoe UI, US Letter, 1in margins, body 11pt.
H1 15pt / H2 13pt navy 0F1E2E with pale-blue rule. Accents mint 7FD4C4, pale blue B9CBDD.
Tables: navy header row, zebra body rows.

Markdown subset supported (same contract as md2docx.ps1):
  first '# '        -> cover banner title
  '**K:** v · ...'  -> grey control block (immediately under title)
  '## ' / '### '    -> headings
  '| a | b |'       -> table (first row = header; '---' separator row skipped)
  '> text'          -> amber callout
  '- ' / '* '       -> bullet
  '1. '             -> numbered
  '---'             -> horizontal rule
  trailing '*text*' -> muted footnote
"""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = "0F1E2E"
MINT = "7FD4C4"
PALE = "B9CBDD"
AMBER = "B5761F"
AMBER_BG = "FDF6E7"
GREY_BG = "F1F4F7"
ZEBRA = "F7F9FB"
MUTED = "5B6B7B"
FONT = "Segoe UI"


def _shade(el, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    el.append(sh)


def _border(p, edge="bottom", color=PALE, sz=8, space=4):
    pPr = p._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    e = OxmlElement("w:" + edge)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    pbdr.append(e)


def _para_shade(p, fill):
    _shade(p._p.get_or_add_pPr(), fill)


def _cell_shade(cell, fill):
    _shade(cell._tc.get_or_add_tcPr(), fill)


def _no_autofit(table):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def _rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    _border(p, "bottom", PALE, 6, 1)


_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def _runs(p, text, size=11, color=None, bold=False, italic=False):
    """Write text into paragraph p, honouring **bold**, *italic* and `code`."""
    for part in _INLINE.split(text):
        if not part:
            continue
        b, i = bold, italic
        t = part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            t, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            t, i = part[1:-1], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            t = part[1:-1]
        r = p.add_run(t)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = b
        r.italic = i
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def _style_base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.12
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)


def _footer(doc, text):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, text, size=8, color=MUTED)


def _table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    _no_autofit(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _cell_shade(c, NAVY)
        c.paragraphs[0].text = ""
        _runs(c.paragraphs[0], h, size=9.5, color="FFFFFF", bold=True)
    for n, row in enumerate(body):
        cells = t.add_row().cells
        for i in range(len(header)):
            c = cells[i]
            if n % 2 == 1:
                _cell_shade(c, ZEBRA)
            c.paragraphs[0].text = ""
            _runs(c.paragraphs[0], row[i] if i < len(row) else "", size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def build(md, out, eyebrow="", classification=""):
    doc = Document()
    _style_base(doc)
    lines = md.split("\n")
    i, first_h1 = 0, True
    pending_table = []

    def flush_table():
        nonlocal pending_table
        if pending_table:
            _table(doc, pending_table)
            pending_table = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()

        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                pending_table.append(cells)
            i += 1
            continue
        flush_table()

        if not s:
            i += 1
            continue

        if s == "---":
            _rule(doc)
            i += 1
            continue

        if s.startswith("# "):
            if first_h1:
                first_h1 = False
                if eyebrow:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(2)
                    _runs(p, eyebrow.upper(), size=8.5, color=MINT, bold=True)
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                _runs(p, s[2:], size=22, color=NAVY, bold=True)
                _border(p, "bottom", MINT, 12, 6)
                if classification:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(10)
                    _runs(p, classification, size=8.5, color=AMBER, bold=True)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                _runs(p, s[2:], size=15, color=NAVY, bold=True)
                _border(p, "bottom", PALE, 8, 4)
            i += 1
            continue

        if s.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(13)
            p.paragraph_format.space_after = Pt(4)
            _runs(p, s[3:], size=13, color=NAVY, bold=True)
            _border(p, "bottom", PALE, 6, 3)
            i += 1
            continue

        if s.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            _runs(p, s[4:], size=11, color=NAVY, bold=True)
            i += 1
            continue

        if s.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            _runs(p, s[5:], size=10, color=MUTED, bold=True)
            i += 1
            continue

        if s.startswith("> "):
            block = []
            while i < len(lines) and lines[i].strip().startswith("> "):
                block.append(lines[i].strip()[2:])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.12)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            _para_shade(p, AMBER_BG)
            _border(p, "left", AMBER, 18, 8)
            _runs(p, " ".join(block), size=10.5, color="4A3A16")
            continue

        # control block: **K:** v · **K:** v
        if s.startswith("**") and "·" in s and s.count("**") >= 2:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(12)
            _para_shade(p, GREY_BG)
            _runs(p, s, size=9.5, color="243B52")
            i += 1
            continue

        m = re.match(r"^(\s*)[-*]\s+(.*)$", raw)
        if m:
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(3)
            _runs(p, m.group(2), size=10.5)
            i += 1
            continue

        m = re.match(r"^\s*\d+\.\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            _runs(p, m.group(1), size=10.5)
            i += 1
            continue

        # trailing muted footnote: whole line italic
        if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            _runs(p, s[1:-1], size=9, color=MUTED, italic=True)
            i += 1
            continue

        p = doc.add_paragraph()
        _runs(p, s, size=11)
        i += 1

    flush_table()
    if classification:
        _footer(doc, classification)
    doc.save(out)
    return out


if __name__ == "__main__":
    import sys
    md = open(sys.argv[1]).read()
    eb = sys.argv[3] if len(sys.argv) > 3 else ""
    cl = sys.argv[4] if len(sys.argv) > 4 else ""
    print(build(md, sys.argv[2], eb, cl))
